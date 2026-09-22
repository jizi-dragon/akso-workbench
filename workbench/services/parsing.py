"""附件解析：把用户交上来的文件变成**模型能读的文本 + 可复算的坐标**。

## 这个模块的产出契约（决定了 G3「能正确完成任务」能不能验）

解析不只是"抽出字符串"，而是三件东西一起给：

| 产出 | 是什么 | 谁用 |
|---|---|---|
| `text` | 正文（有硬上限，超了截断并**明说截到哪**） | 内联进上下文 / 工具按范围取 |
| `locator` | **坐标目录**：页 / 表 / 区域 / 段落号 | 模型回答里的"第 7 页""明细表 B4"要能**被工具复算出来** |
| `note` | 解析过程里"不完美"的地方（截断、跳过的对象、扫描件无文本…） | 如实告诉模型与用户，不许让它以为读全了 |

`note` 那一条是本仓反复吃亏的地方（"接口回成功、回读才发现空"）：解析失败**不能**表现为
"一份空文件"，那会让模型理直气壮地回答"文件里没有这个内容"。

## 为什么按扩展名分流 + 内容嗅探复核

扩展名可以骗人（`.docx` 里塞个 zip、`.pdf` 里塞文本）。所以：
① 扩展名决定**用哪个解析器**；② 魔数（magic）复核它**确实是那个东西**，不是就拒收。
两者都不通过 ⇒ 拒收（Fail-Closed），不做"猜一个解析器试试"。

## 诚实边界（写在代码里，不写在承诺里）

- PDF 抽的是**文本流**：双栏/表格会串行（要版式级正确得换引擎，本模块不承诺）；
- 扫描件（图片型 PDF）抽不出文本 ⇒ `note` 里明说"这可能是扫描件"，而不是回空串；
- 图片**不解析**（走模型多模态，见 `routes_chat._model_history`）；
- xlsx 只读**值**，不读公式结果缓存以外的样式/图表。
"""

from __future__ import annotations

import csv
import io
import json
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .. import config

#: 扩展名 → (kind, 人类可读类型名)。**白名单**：表里没有的一律拒收。
SUPPORTED: dict[str, tuple[str, str]] = {
    ".md": ("text", "Markdown"),
    ".markdown": ("text", "Markdown"),
    ".txt": ("text", "文本"),
    ".csv": ("sheet", "CSV 表格"),
    ".docx": ("document", "Word 文档"),
    ".xlsx": ("workbook", "Excel 工作簿"),
    ".pdf": ("pdf", "PDF 文档"),
    ".png": ("image", "图片"),
    ".jpg": ("image", "图片"),
    ".jpeg": ("image", "图片"),
    ".webp": ("image", "图片"),
}

#: 魔数 → 允许的扩展名集合（`None` = 该类型没有稳定魔数，只能靠扩展名）。
#: 校验方向是"**这个字节流像不像它声明的类型**"，不是"能不能猜出它是什么"。
_MAGIC: tuple[tuple[bytes, frozenset[str]], ...] = (
    (b"%PDF-", frozenset({".pdf"})),
    (b"PK\x03\x04", frozenset({".docx", ".xlsx"})),   # OOXML 都是 zip
    (b"\x89PNG\r\n\x1a\n", frozenset({".png"})),
    (b"\xff\xd8\xff", frozenset({".jpg", ".jpeg"})),
)

IMAGE_EXTENSIONS = frozenset({".png", ".jpg", ".jpeg", ".webp"})
WEBP_MAGIC = (b"RIFF", b"WEBP")


@dataclass
class Parsed:
    kind: str = "text"
    text: str = ""
    locator: dict[str, Any] = field(default_factory=dict)
    note: str = ""
    truncated: bool = False


def extension_of(name: str) -> str:
    return Path(str(name or "")).suffix.lower()


def type_label(ext: str) -> str:
    return SUPPORTED.get(ext, ("", "未知类型"))[1]


def sniff_ok(data: bytes, ext: str) -> bool:
    """魔数复核：**只认能确认的类型**；`txt/md/csv` 没有稳定魔数 ⇒ 按扩展名放行。"""
    if ext in IMAGE_EXTENSIONS:
        if ext == ".webp":
            return data[:4] == WEBP_MAGIC[0] and data[8:12] == WEBP_MAGIC[1]
        return any(data.startswith(magic) and ext in allowed
                   for magic, allowed in _MAGIC if ext in allowed)
    if ext in (".txt", ".md", ".markdown", ".csv"):
        return True
    for magic, allowed in _MAGIC:
        if ext in allowed:
            return data.startswith(magic)
    return False


# ---------------------------------------------------------------- 各类型的解析


class _Builder:
    """按"单元"拼正文，并**同时记下每个单元在正文里的字符区间**。

    为什么非要区间不可：G3 要求"模型说第 7 页，工具读第 7 页是同一件事"。
    没有区间，`attachment_read(page=7)` 就只能靠猜（重新解析一遍、或按空行切），
    而"猜出来的第 7 页"与"模型看到的第 7 页"可能不是同一段——那是最难查的一类错。

    区间是**字符偏移**（不是行号）：正文是拼出来的，行号会随拼接方式漂，偏移不会。
    """

    def __init__(self) -> None:
        self.parts: list[str] = []
        self.offset = 0
        self.units: list[dict[str, Any]] = []

    def add(self, text: str, *, record: bool = True, **unit: Any) -> None:
        """追加一段，并把它的区间记成 `unit`。

        `record=False` 用于"这一段的区间由外层单元统一覆盖"的行（如表格的每一行）——
        逐行记账会造出一堆没有区分键的假单元，而它们会被读侧当成"可选坐标"。
        """
        if not text:
            return
        self.parts.append(text)
        if record:
            self.units.append({**unit, "offset": self.offset, "chars": len(text)})
        self.offset += len(text)

    @property
    def text(self) -> str:
        return "".join(self.parts)


def _decode_text(data: bytes) -> str:
    """文本解码：UTF-8 优先，失败退 GBK（中文环境里 .csv/.txt 常见），再退替换。"""
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _clip(text: str, limit: int) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[:limit], True


def _parse_plain(text: str, *, limit: int) -> Parsed:
    body, truncated = _clip(text, limit)
    # 行数按 `splitlines()` 数：末尾那个换行不算"多了一行"（否则"第 4 行"会指向空白）
    lines = len(text.splitlines())
    locator = {"lines": lines, "head": [ln for ln in text.splitlines()[:5]]}
    note = f"正文 {len(text)} 字，已截断到 {limit} 字" if truncated else ""
    return Parsed(kind="text", text=body, locator=locator, note=note, truncated=truncated)


def _parse_csv(text: str, *, limit: int) -> Parsed:
    """CSV：**一行 = 一条记录**，正文按行拼 ⇒ 行号就是可复算的坐标（读侧按行取）。"""
    rows = list(csv.reader(io.StringIO(text)))
    header = rows[0] if rows else []
    body = "\n".join(" | ".join(cell for cell in row) for row in rows)
    clipped, truncated = _clip(body, limit)
    locator = {"rows": len(rows), "cols": len(header), "header": header[:12]}
    note = f"共 {len(rows)} 行，正文已截断到 {limit} 字" if truncated else ""
    return Parsed(kind="sheet", text=clipped, locator=locator, note=note, truncated=truncated)


def _parse_docx(path: Path, *, limit: int) -> Parsed:
    import docx  # 延迟导入：没装依赖时不该让整个模块 import 失败

    document = docx.Document(str(path))
    builder = _Builder()
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style = str(getattr(paragraph.style, "name", "") or "")
        chunk = f"[{style}] {text}" if style.startswith("Heading") else text
        builder.add(chunk + "\n", unit="p", p=index, style=style, head=text[:40])
    for t_index, table in enumerate(document.tables, start=1):
        rows = [[(cell.text or "").strip() for cell in row.cells] for row in table.rows]
        start = builder.offset
        builder.add(f"\n[表 {t_index}]\n", record=False)
        for row in rows:
            builder.add(" | ".join(row) + "\n", record=False)
        # 整个表记成**一个**单元（含表头行）——逐行记账会造出一堆没有区分键的假单元
        builder.units.append({"table": t_index, "offset": start,
                              "chars": builder.offset - start, "rows": len(rows),
                              "cols": len(rows[0]) if rows else 0,
                              "header": rows[0][:12] if rows else []})
    full = builder.text
    body, truncated = _clip(full, limit)
    locator: dict[str, Any] = {"units": [u for u in builder.units if "p" in u][:200],
                               "tables": [u for u in builder.units if "table" in u]}
    if not locator["units"]:
        locator = {"tables": locator["tables"]}
    note = f"正文已截断到 {limit} 字" if truncated else ""
    return Parsed(kind="document", text=body, locator=locator, note=note, truncated=truncated)


def _parse_xlsx(path: Path, *, limit: int) -> Parsed:
    import openpyxl

    book = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    builder = _Builder()
    sheets: list[dict[str, Any]] = []
    try:
        for sheet in book.worksheets:
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(["" if cell is None else str(cell) for cell in row])
            start = builder.offset
            builder.add(f"\n[工作表 {sheet.title}]（{len(rows)} 行 × {sheet.max_column or 0} 列）\n")
            for index, row in enumerate(rows, start=1):
                builder.add(f"{index}: " + " | ".join(row) + "\n")
            sheets.append({"sheet": sheet.title, "offset": start,
                           "chars": builder.offset - start, "rows": len(rows),
                           "cols": sheet.max_column or 0,
                           "header": rows[0][:12] if rows else []})
            if builder.offset > limit * 2:
                break
    finally:
        book.close()
    body, truncated = _clip(builder.text, limit)
    notes: list[str] = []
    if truncated:
        notes.append(f"正文已截断到 {limit} 字")
    if len(sheets) < (book.sheetnames and len(book.sheetnames) or 0):
        notes.append("有工作表因超出上限未展开")
    return Parsed(kind="workbook", text=body, locator={"sheets": sheets},
                  note="；".join(notes), truncated=truncated)


def _parse_pdf(path: Path, *, limit: int) -> Parsed:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    builder = _Builder()
    empty_pages = 0
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:  # noqa: BLE001 —— 单页坏掉不该让整份文件读不出来
            text = ""
        if not text:
            empty_pages += 1
        builder.add(f"\n[第 {number} 页]\n{text}\n", unit="page", page=number,
                    head=" ".join(text.split())[:60])
    pages = builder.units
    body, truncated = _clip(builder.text, limit)
    notes: list[str] = []
    if truncated:
        notes.append(f"正文已截断到 {limit} 字")
    if pages and empty_pages == len(pages):
        # **整份都没有文本** ⇒ 正文回空，而不是留着 "[第 N 页]" 这类骨架。
        # 留着骨架的后果很具体：`text_chars` 会是一个非零的数字（"看起来有内容"），
        # 于是上游的内联块与状态判断都会以为"解析成功了、只是内容少"。
        body = ""
        truncated = False
        notes.append("**这一份没有可抽取的文本**：很可能是扫描件/图片型 PDF"
                     "（本工作台不做 OCR）——请让用户改传原始文件或截图，"
                     "**不要**据它回答文档内容")
    elif empty_pages:
        notes.append(f"有 {empty_pages} 页没有文本（可能是图）")
    return Parsed(kind="pdf", text=body, locator={"pages": pages, "pageCount": len(pages)},
                  note="；".join(notes), truncated=truncated)


# ---------------------------------------------------------------- 入口


def parse_file(path: Path, *, name: str = "", limit: int | None = None) -> Parsed:
    """按扩展名解析。**任何异常都转成"解析失败"的 Parsed**（带 note），不向上抛。

    为什么不抛：调用方（上传接口）已经拿到并落盘了文件——解析失败不该让整个上传失败，
    但也**绝不能**伪装成"文件是空的"（那正是本仓最贵的一类缺陷）。
    """
    budget = int(limit or config.ATTACHMENTS_MAX_TEXT_CHARS)
    ext = extension_of(name or path.name)
    kind = SUPPORTED.get(ext, ("text", ""))[0]
    if kind == "image":
        return Parsed(kind="image", text="", locator={},
                      note="图片不解析文本，直接交给模型看图（多模态）")
    try:
        if ext in (".txt", ".md", ".markdown"):
            return _parse_plain(_decode_text(path.read_bytes()), limit=budget)
        if ext == ".csv":
            return _parse_csv(_decode_text(path.read_bytes()), limit=budget)
        if ext == ".docx":
            return _parse_docx(path, limit=budget)
        if ext == ".xlsx":
            return _parse_xlsx(path, limit=budget)
        if ext == ".pdf":
            return _parse_pdf(path, limit=budget)
    except Exception as exc:  # noqa: BLE001 —— 见 docstring
        return Parsed(kind=kind or "text", text="", locator={},
                      note=f"解析失败（{type(exc).__name__}: {exc}）⇒ "
                           "**不要**据它推断内容，请让用户换一种格式重传或说明文件结构")
    return Parsed(kind=kind or "text", text="", locator={},
                  note=f"不支持的类型：{ext or '（无扩展名）'}")


def locator_json(parsed: Parsed) -> str:
    try:
        return json.dumps(parsed.locator, ensure_ascii=False)[:20_000]
    except (TypeError, ValueError):
        return "{}"


def zip_is_intact(path: Path) -> bool:
    """OOXML（docx/xlsx）额外体检：zip 能打开且含 `[Content_Types].xml`。

    为什么单独一条：魔数只说明"它是个 zip"。把改名成 .docx 的普通 zip 交给
    `python-docx`，抛出来的错误形态是"缺某个 xml"——用户看到的是"解析失败"，
    而真实原因是"这根本不是 Word 文件"。这一条把两者分开。
    """
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
    except (zipfile.BadZipFile, OSError):
        return False
    return "[Content_Types].xml" in names
