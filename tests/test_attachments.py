"""附件能力验收：上传/落盘/解析/**真的进到模型请求**（2026-09-22，用户下达的核心能力）。

> 用户原话：「核心 agent 能力增强，未来，用户可能会直接传一张图片，或者是传文件
> （word，pdf，markdown，excel）给 agent，让 agent 去读取。这样的话我就得保障：
> ① 能够让用户正常的输入。② 能够让 agent 真实的获取到相关文件。
> ③ 能够正确的按照知识链路，完成相关的任务。」

本文件钉的是 **② 与 ③ 的地基**（① 的 UI 部分在 `test_user_confirmation.py` 的静态契约里）：

| 组 | 钉什么 | 为什么非钉不可 |
|---|---|---|
| A 落盘 | 字节数 / sha256 / 路径不越界 / 删除真删 | "存进去了"与"存对了"是两件事；路径必须不接受用户输入 |
| B 拒收 | 类型白名单 + 魔数 + 空文件 + 超限 + 假 docx | 拒收路径**最容易静默放行**（改个后缀就过） |
| C 解析 | 每类型都能抽到正文 + **坐标**（页/表/行） | 坐标是"回答能指到出处"的全部依据 |
| D 进模型 | 图片 → `image_url`；文档 → 文本块；**只有最新一轮**展开 | 这条是"agent 真的拿到了文件"，不是"接口回了 200" |
| E 按需取 | `attachment_read` 按坐标取到的，与解析时记的区间**逐字一致** | 判据可复算，模型说第 7 页 = 工具读第 7 页 |
| F 隔离 | 别的会话的附件读不到；token 估算把图片算进去 | 数据串台 + 账单失控都是静默的 |
"""

from __future__ import annotations

import asyncio
import json
import struct
import zlib
from pathlib import Path

import pytest

from workbench import db
from workbench import plugin_api as pa
from workbench.services import attachments as store


@pytest.fixture(autouse=True)
def _db():
    db.connect()
    yield


# ---------------------------------------------------------------- 造测试文件


def tiny_png() -> bytes:
    """8×8 红底 PNG（手搓，不引依赖）。"""
    width = height = 8
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))

    def chunk(tag: bytes, payload: bytes) -> bytes:
        return (struct.pack(">I", len(payload)) + tag + payload
                + struct.pack(">I", zlib.crc32(tag + payload) & 0xFFFFFFFF))

    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw))
            + chunk(b"IEND", b""))


def make_docx(tmp_path: Path) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("培训记录规范", level=1)
    document.add_paragraph("第一条：培训结束后 3 个工作日内登记。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "要求"
    table.cell(1, 0).text = "签到表"
    table.cell(1, 1).text = "需本人签名"
    path = tmp_path / "spec.docx"
    document.save(path)
    return path.read_bytes()


def make_xlsx(tmp_path: Path) -> bytes:
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "明细"
    sheet.append(["部门", "人数", "费用"])
    for row in (("生产部", 12, 3600), ("质量部", 5, 1500)):
        sheet.append(row)
    path = tmp_path / "cost.xlsx"
    book.save(path)
    return path.read_bytes()


def make_pdf(text: str = "AKSO-42 training record") -> bytes:
    """手搓一页 PDF（含可抽取文本）。**不引第三方写库**：pypdf 只负责读。"""
    content = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 200]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length " + str(len(content)).encode() + b">>stream\n" + content + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj".encode() + body + b"endobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (f"trailer<</Size {len(objects) + 1}/Root 1 0 R>>\nstartxref\n"
            f"{xref_at}\n%%EOF\n").encode()
    return bytes(out)


def upload(conversation: str, name: str, data: bytes) -> dict:
    return store.store(conversation_id=conversation, name=name, data=data)


# ---------------------------------------------------------------- A 落盘


def test_store_writes_bytes_and_metadata_truthfully() -> None:
    """落盘的字节、sha256、大小必须与上传的一致——**这是"agent 真的拿到了文件"的地基**。"""
    import hashlib

    data = "# 标题\n\n正文甲乙丙\n".encode()
    row = upload("conv-a", "说明.md", data)

    assert row["bytes"] == len(data)
    assert row["sha256"] == hashlib.sha256(data).hexdigest()
    assert row["kind"] == "text"
    assert row["status"] == "parsed"
    assert Path(row["stored_name"]).name == row["stored_name"], "落盘名里混进了路径"

    path = store.raw_path(row)
    assert path.read_bytes() == data, "磁盘上的内容与上传的不一致"
    assert store.parsed_text(row) == "# 标题\n\n正文甲乙丙\n"

    assert store.delete(row["id"]) is True
    assert not path.exists(), "删了记录但文件还在"
    assert store.get(row["id"]) is None


# ---------------------------------------------------------------- B 拒收


@pytest.mark.parametrize("name,data,match", [
    ("病毒.exe", b"MZ\x90\x00", "不支持的文件类型"),
    ("空.txt", b"", "空的"),
    ("改名.docx", b"PK\x03\x04not-a-real-docx", "有效的 .docx"),
    ("假.png", b"not a png at all", "内容与扩展名不符"),
])
def test_upload_rejects_bad_inputs(name: str, data: bytes, match: str) -> None:
    """拒收路径必须**真的拒**：类型白名单、空文件、魔数复核（改后缀骗不过去）。"""
    with pytest.raises(ValueError, match=match):
        upload("conv-b", name, data)


def test_upload_rejects_oversize(monkeypatch) -> None:
    from workbench import config

    monkeypatch.setattr(config, "ATTACHMENTS_MAX_BYTES", 16)
    with pytest.raises(ValueError, match="文件过大"):
        upload("conv-b", "大.txt", b"x" * 64)


def test_upload_rejects_a_zip_renamed_to_docx() -> None:
    """真 zip 但**不是** OOXML：魔数过得了，`[Content_Types].xml` 体检过不了。

    这一条是"错误信息要指对原因"的判据：模型/用户看到的应当是"这不是有效的 docx"，
    而不是 python-docx 抛出来的"缺某个 xml"。
    """
    import io
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "hi")
    with pytest.raises(ValueError, match="有效的 .docx"):
        upload("conv-b", "伪装.docx", buffer.getvalue())


# ---------------------------------------------------------------- C 解析 + 坐标


def test_parse_markdown_keeps_line_coordinates() -> None:
    row = upload("conv-c", "笔记.md", b"# A\nline two\nline three\n")
    assert row["status"] == "parsed"
    locator = store.locator(row)
    assert locator["lines"] == 3
    assert "line two" in store.parsed_text(row)


def test_parse_csv_keeps_rows_and_header() -> None:
    row = upload("conv-c", "名单.csv", "姓名,部门\n张三,生产部\n李四,质量部\n".encode())
    locator = store.locator(row)
    assert locator["rows"] == 3 and locator["header"] == ["姓名", "部门"]
    assert "张三" in store.parsed_text(row)


def test_parse_docx_keeps_paragraph_and_table_coordinates(tmp_path: Path) -> None:
    row = upload("conv-c", "规范.docx", make_docx(tmp_path))
    text = store.parsed_text(row)
    assert row["kind"] == "document" and row["status"] == "parsed"
    assert "培训记录规范" in text and "签到表" in text, text[:200]
    locator = store.locator(row)
    assert locator["units"] and locator["tables"], locator
    table = locator["tables"][0]
    assert table["rows"] == 2 and table["header"][:1] == ["项目"]
    # 坐标必须**能切回原文**（否则"第几段"只是装饰）
    assert "签到表" in text[table["offset"]:table["offset"] + table["chars"]]


def test_parse_xlsx_keeps_sheet_coordinates(tmp_path: Path) -> None:
    row = upload("conv-c", "费用.xlsx", make_xlsx(tmp_path))
    text = store.parsed_text(row)
    assert row["kind"] == "workbook"
    assert "生产部" in text and "3600" in text
    sheet = store.locator(row)["sheets"][0]
    assert sheet["sheet"] == "明细" and sheet["rows"] == 3
    assert sheet["header"] == ["部门", "人数", "费用"]
    assert "3600" in text[sheet["offset"]:sheet["offset"] + sheet["chars"]]


def test_parse_pdf_keeps_page_coordinates() -> None:
    row = upload("conv-c", "报告.pdf", make_pdf("AKSO-42 training record"))
    text = store.parsed_text(row)
    assert row["kind"] == "pdf", row["parse_note"]
    locator = store.locator(row)
    assert locator["pageCount"] == 1
    page = locator["pages"][0]
    assert page["page"] == 1 and page["offset"] == 0
    assert "AKSO-42" in text[page["offset"]:page["offset"] + page["chars"]], text[:200]


def test_scanned_pdf_says_so_instead_of_pretending_to_be_empty() -> None:
    """**空解析必须出声**：没有文本的 PDF 要说"可能是扫描件、别据此作答"。

    同时钉住"正文回空"：留着 `[第 N 页]` 骨架的话 `text_chars` 会是个非零数字，
    上游会以为"解析成功了、只是内容少"——那正是本仓最贵的一类静默失真。
    """
    row = upload("conv-c", "扫描件.pdf", make_pdf(""))  # 空内容流 ⇒ 抽不到文本
    note = str(row["parse_note"])
    assert row["text_chars"] == 0, row
    assert "没有可抽取的文本" in note and "扫描件" in note, note


# ---------------------------------------------------------------- D 真的进模型请求


def _history(conversation: str, rows: list[tuple[str, str]]) -> list[dict]:
    from workbench.api.routes_chat import _model_history

    return _model_history([
        {"role": role, "content": content} for role, content in rows
    ])


def test_image_attachment_becomes_an_image_part_in_the_newest_turn() -> None:
    """★ "agent 真的拿到了图片"的判据：**请求体里是 image_url**，不是一句文件名。

    顺带钉住那句随图片一起发的提示——它**是承重的**：模型会先调 `attachment_list`，
    看到图片 `textChars=0` 就推断"文件空的/坏的"并**拒绝看图**（2026-09-22 真机实录：
    「看不出颜色——这个文件只有 168 字节、textChars=0，像是空文件或损坏的图」，
    而像素明明就在上下文里）。所以提示里必须明说"像素在下面、别拿元数据当结论"。
    """
    row = upload("conv-d", "截图.png", tiny_png())
    history = _history("conv-d", [
        ("user", "看看这张图 " + store.reference_marker(row)),
    ])
    content = history[-1]["content"]
    assert isinstance(content, list), content
    kinds = [part["type"] for part in content]
    assert "image_url" in kinds, kinds
    url = next(p for p in content if p["type"] == "image_url")["image_url"]["url"]
    assert url.startswith("data:image/png;base64,"), url[:40]
    note = "\n".join(p.get("text", "") for p in content)
    assert "见图" in note and "图片" in note, note
    assert "能不能读" not in note and "textChars" not in note, (
        "提示里不要提起'能不能读'——2026-09-22 实测：带这层解释的那版提示之后，"
        "模型当轮就回「看不出颜色……像是空文件或损坏的图」（把念头递给了它）")


def test_document_attachment_inlines_text_and_coordinates(tmp_path: Path) -> None:
    row = upload("conv-d", "规范.docx", make_docx(tmp_path))
    history = _history("conv-d", [("user", "按这份规范检查 " + store.reference_marker(row))])
    text = "\n".join(p.get("text", "") for p in history[-1]["content"])
    assert "签到表" in text, "正文没有进上下文"
    assert "坐标目录" in text, "没给坐标 ⇒ 模型无法指到出处"
    assert row["id"] in text


def test_older_turns_do_not_resend_the_file(tmp_path: Path) -> None:
    """★ 用户拍板的 A4：**只有最新一轮**展开，历史轮只留一句引用。

    判据同时钉两件事：① 旧轮没有 image_url/正文；② **最新轮仍然有**（不是干脆不展开了）。
    """
    first = upload("conv-e", "旧图.png", tiny_png())
    second = upload("conv-e", "新图.png", tiny_png())
    history = _history("conv-e", [
        ("user", "第一张 " + store.reference_marker(first)),
        ("assistant", "收到"),
        ("user", "第二张 " + store.reference_marker(second)),
    ])
    old, newest = history[0]["content"], history[-1]["content"]
    old_parts = old if isinstance(old, list) else [{"type": "text", "text": old}]
    old_text = "\n".join(p.get("text", "") for p in old_parts)
    assert not any(p["type"] == "image_url" for p in old_parts), (
        "历史轮重发了图片（每轮都要重发一遍 ⇒ 账单与上下文一起爆）")
    assert "你在那一轮看过它的像素" in old_text and first["id"] in old_text, old_text
    assert isinstance(newest, list) and any(p["type"] == "image_url" for p in newest)
    assert second["id"] in "\n".join(p.get("text", "") for p in newest)


def test_older_image_note_must_not_tell_the_model_it_never_saw_it() -> None:
    """旧轮次的图片引用话术：**不能说成"内容没有发过"**（2026-09-22 真机联调抓到）。

    现场实录：模型答对了纯蓝色块与黑方块，第三轮却**主动撤回**——
    「前两轮关于色块和形状的图片，我并没有真的读取附件内容，回答是瞎猜的」。
    原因是旧轮那句与文档共用的话术写着「已交过你、**正文**没有重复发送、要看内容用
    `attachment_read`」：图片没有"正文"，而 `attachment_read` 也给不了像素
    ⇒ 模型读到的是"我从来没见过这张图"，于是把**自己正确的感知**否掉了。

    判据（两条，都是可执行的字符串断言）：① 图片的旧轮提示必须说清"**你当时看过**、
    现在看不到细节、要再看请让用户重贴"；② 不许把 `attachment_read` 说成能看图，
    也不许出现「正文没有重复发送」这种对图片不成立的措辞。
    """
    att = upload("conv-note", "旧图.png", tiny_png())
    history = _history("conv-note", [
        ("user", "第一张 " + store.reference_marker(att)),
        ("assistant", "蓝色"),
        ("user", "再看看这个表格"),
    ])
    old = history[0]["content"]
    text = "\n".join(p.get("text", "") for p in (old if isinstance(old, list) else []))

    assert "看过它的像素" in text, text
    assert "重新贴" in text, text
    assert "正文没有重复发送" not in text, (
        "这句话对图片不成立（图片没有正文），模型会据此否掉自己刚才的答案")
    assert "attachment_read" not in text or "给不了像素" in text, (
        "不许把 attachment_read 说成能看图——它只回文件名与类型")



def test_estimate_tokens_counts_the_image() -> None:
    """图片不折算 token ⇒ 上下文阈值形同虚设（图片会一直堆着不被折叠）。"""
    from workbench.harness.context import IMAGE_TOKEN_ESTIMATE, estimate_content_tokens

    parts = [{"type": "text", "text": "看图"},
             {"type": "image_url", "image_url": {"url": "data:image/png;base64,AAAA"}}]
    assert estimate_content_tokens(parts) >= IMAGE_TOKEN_ESTIMATE
    assert estimate_content_tokens("看图") < IMAGE_TOKEN_ESTIMATE


# ---------------------------------------------------------------- E 按坐标取（可复算）


def _call_tool(name: str, args: dict, conversation: str) -> dict:
    from workbench.harness.tools import default_registry

    registry = default_registry()
    invocation = asyncio.run(registry.execute(
        name, args, pa.ToolContext(conversation_id=conversation, source="test")))
    assert invocation.status is pa.ToolStatus.OK, invocation.content
    return json.loads(invocation.content)


def test_attachment_read_by_page_matches_the_recorded_range(tmp_path: Path) -> None:
    """★ G3 的判据：**工具按坐标取到的，与解析时记的区间逐字一致**。

    这条是"模型说第 7 页、工具读第 7 页是同一件事"的可执行版本。
    """
    row = upload("conv-f", "报告.pdf", make_pdf("AKSO-42 training record"))
    page = store.locator(row)["pages"][0]
    payload = _call_tool("attachment_read",
                         {"attachment_id": row["id"], "unit": "page", "index": 1}, "conv-f")
    assert payload["ok"] is True
    assert "AKSO-42" in payload["text"]
    assert payload["text"] == store.parsed_text(row)[page["offset"]:page["offset"] + page["chars"]]


def test_attachment_read_by_sheet_and_missing_targets(tmp_path: Path) -> None:
    row = upload("conv-f", "费用.xlsx", make_xlsx(tmp_path))
    ok = _call_tool("attachment_read",
                    {"attachment_id": row["id"], "unit": "sheet", "name": "明细"}, "conv-f")
    assert "生产部" in ok["text"] and ok["range"] == "工作表 '明细'"

    miss = _call_tool("attachment_read",
                      {"attachment_id": row["id"], "unit": "sheet", "name": "不存在的表"}, "conv-f")
    assert miss["ok"] is False and "没有" in miss["range"], miss


def test_attachment_list_shows_only_this_conversation() -> None:
    upload("conv-g1", "甲.md", b"a\n")
    upload("conv-g2", "乙.md", b"b\n")
    payload = _call_tool("attachment_list", {}, "conv-g1")
    names = [a["name"] for a in payload["attachments"]]
    assert names == ["甲.md"], names


def test_attachment_list_tells_the_model_images_have_no_text_but_are_visible() -> None:
    """`attachment_list` 里图片那条必须自带"没有文本 ≠ 读不出"（2026-09-22 真机实录）。

    模型看到 `textChars: 0` 会推断"这个是空文件/损坏的图"，于是**拒绝看图**——
    实录回答「看不出颜色——这个文件只有 168 字节、textChars=0，像是空文件或损坏的图」，
    而同一会话的下一问它又答对了"正方形"（说明像素确实在）。**工具回什么，模型就信什么**，
    所以这句话必须由工具自己说清，不能指望模型去猜。
    """
    upload("conv-g3", "证据.png", tiny_png())
    row = _call_tool("attachment_list", {}, "conv-g3")["attachments"][0]
    assert row["kind"] == "image" and row["textChars"] == 0
    assert "像素" in row["note"] and "不要用本工具判断" in row["note"], row["note"]


def test_cross_conversation_read_is_refused() -> None:
    """数据串台是静默的：模型会答得头头是道，而内容来自另一段对话。"""
    row = upload("conv-h1", "报价单.md", "单价 100\n".encode())
    from workbench.harness.tools import default_registry

    registry = default_registry()
    invocation = asyncio.run(registry.execute(
        "attachment_read", {"attachment_id": row["id"]},
        pa.ToolContext(conversation_id="conv-h2", source="test")))
    assert invocation.status is pa.ToolStatus.ERROR
    assert "另一个会话" in invocation.content, invocation.content


def test_pending_upload_is_bound_on_first_send() -> None:
    """★ 新会话的第一条消息之前上传 ⇒ 落在"未绑定"哨兵下 ⇒ **发言时必须绑上**。

    这条是 2026-09-22 **真机联调**抓到的缺陷的判据：绑定条件原本只判空串，
    而哨兵 `_pending` 是**真值** ⇒ 永远不绑 ⇒ 附件留在哨兵目录里，
    会话里的 `attachment_list` 看不到它、`attachment_read` 还回一句"属于另一个会话"。
    而模型的表现是**合理但错误**的："这次不能看：该附件属于另一段会话"。
    """
    from workbench.services.attachments import PENDING_CONVERSATION

    row = upload(PENDING_CONVERSATION, "待绑.csv", "部门,人数\n生产部,12\n".encode())
    assert row["conversation_id"] == PENDING_CONVERSATION

    bound = store.attach_to_conversation(row["id"], "conv-real")
    assert bound is not None and bound["conversation_id"] == "conv-real", bound

    # ★ 文件也必须跟着搬过去：正文路径是**从 conversation_id 推出来的**，
    # 只改库不搬文件 ⇒ 读的时候去新目录里找一个还躺在 _pending/ 下的文件
    # （真机实测的形态：图片问颜色回"读取失败：No such file or directory"）
    assert store.raw_path(bound).is_file(), f"原件没搬过去：{store.raw_path(bound)}"
    assert store.parsed_text(bound), "解析产物没搬过去（按范围取正文会读到空）"

    # 绑上之后，会话内的工具就该看得见它了
    payload = _call_tool("attachment_list", {}, "conv-real")
    assert [a["name"] for a in payload["attachments"]] == ["待绑.csv"], payload
    read = _call_tool("attachment_read",
                      {"attachment_id": row["id"], "unit": "line", "start": 1}, "conv-real")
    assert read["ok"] is True and "生产部" in read["text"], read


def test_binding_never_steals_an_attachment_from_another_conversation() -> None:
    """绑定只允许"未绑定 → 已绑定"：已属于别的会话的附件不许被改嫁。"""
    row = upload("conv-owner", "我的.md", b"x\n")
    again = store.attach_to_conversation(row["id"], "conv-thief")
    assert again is not None and again["conversation_id"] == "conv-owner", again


def test_the_newest_message_goes_through_the_choke_point() -> None:
    """★ **这一轮**的用户消息也必须过扼流点（附件在那里展开）。

    这条钉的是一个**单元测试看不见的**缺陷（2026-09-22 真机联调抓到）：
    真实路径原本是"先取历史 → 过硬点 → 再把这一轮的消息原样追加"，
    于是历史行展开了、**刚交上来的那份反而没有** —— 模型看到的是一个文件名，
    回"没有把像素送过来"，而 `test_image_attachment_becomes_an_image_part_*` 全绿
    （它直接喂的是硬点）。

    所以这条刻意走**真路径**：`_request_history()` 落库 + 过硬点，一步都不省。
    """
    from workbench.api.routes_chat import _request_history
    from workbench.services import storage

    # 真会话（`messages.conversation_id` 有外键 ⇒ 不能随便编一个 id）
    conversation = storage.create_conversation()["id"]
    row = upload(conversation, "图.png", tiny_png())
    history = _request_history(conversation, f"看看这张图\n{store.reference_marker(row)}")
    content = history[-1]["content"]
    assert isinstance(content, list), content
    assert any(p.get("type") == "image_url" for p in content), (
        "这一轮的消息绕过了扼流点 ⇒ 图片没被展开成 image_url")
    assert history[-1]["role"] == "user"
    # 同一段历史再取一次（模拟第二轮请求）：这一轮的图**不该**再被展开
    second = _request_history(conversation, "还有别的吗")
    old = second[-2]["content"]
    old_parts = old if isinstance(old, list) else [{"type": "text", "text": old}]
    assert not any(p.get("type") == "image_url" for p in old_parts), old_parts
